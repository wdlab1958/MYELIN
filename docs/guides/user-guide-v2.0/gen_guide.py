# -*- coding: utf-8 -*-
"""MYELIN 사용자 가이드 v2.0.0 생성기 — docx + xlsx (PDF 는 soffice 로 변환).

출력: docs/guides/user-guide-v2.0/{MYELIN_사용자가이드_v2.0.docx, MYELIN_사용자가이드_표.xlsx, img/*.png}
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(sys.argv[1]).resolve()
OUT = REPO / "docs/guides/user-guide-v2.0"
IMG = OUT / "img"
V1_IMG = REPO / "docs/guides/user-guide-v1.0/img"
P6 = REPO / "docs/reports/P6"

GUIDE_VER = "v2.0.0"
GUIDE_DATE = "2026-10-27"
PRODUCT_VER = "v1.3.1"
ORG = "WDLAB · A3 Security · AI·ML·AX Research Division"
XLSX_NAME = "MYELIN_사용자가이드_표.xlsx"

# Emboss tokens (design/tokens.json)
INK = "2D2A3E"
MUTED = "5F6378"
VIOLET = "7C3AED"
VIOLET_50 = "F3F0FF"
SURFACE_2 = "E6E9F2"
BORDER = "DADEE9"
LIME_INK = "3F6A0A"
CYAN_INK = "0E7490"
MAGENTA_INK = "BE188C"

FONT = "Malgun Gothic"
MONO = "Consolas"

# ----------------------------------------------------------------------------- 표 데이터
# 각 표: (xlsx 시트 이름, 헤더, 행, 열 너비(cm))
TABLES: dict[str, tuple[str, list[str], list[list[str]], list[float]]] = {}


def T(key: str, sheet: str, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    assert all(len(r) == len(headers) for r in rows), key
    TABLES[key] = (sheet, headers, rows, widths)


T("changes", "0 v1.0→v2.0 변경 요약",
  ["영역", "v1.0 (2026-10-18)", "v2.0 (제품 v1.3.1, 2026-10-27)", "근거"],
  [
      ["Review Gate 승격", "Promote 를 바로 누를 수 있었다", "\"내용 읽기\" 로 본문을 연 뒤에만 Promote 가 켜진다(읽음 확인 장치). 같은 이름의 노트가 Domain 폴더에 있으면 409 로 막는다", "정본 Rev 2.1 §7.4"],
      ["Dashboard 지표", "KPI 6종", "KPI 9종 — 재사용 비율·재사용 지연·산출물당 출처가 추가. 각 지표에 달성 / 목표 이탈 배지. Myelination(강도 분포) 카드", "ADR-0012"],
      ["미엘린 강도", "없음", "노트별 강도 0~3 을 myelin ×N 텍스트 라벨로 표시(Dashboard · My Domains 노트 표)", "ADR-0012 · Rev 2.1 §11.4"],
      ["환류(refeed)", "없음", "/compose 산출물에 ## 환류 후보 절이 생기고, /distill <산출물> 이 그 절을 원자 노트 초안으로 되돌린다. /review 가 환류 대기를 보고", "Rev 2.1 §7.3"],
      ["패치 적용", "Apply 즉시 적용", "미리보기(diff · 검사 결과 · digest) → Apply 는 digest 가 일치할 때만(불일치 409). under 구획 append, ## set 허용 키 제한, 처리된 패치는 _patches/applied|rejected/ 로 이동", "ADR-0006 보강"],
      ["기밀 라우팅", "myelin.yaml 의 llm.route 문장", "장치화 — MYELIN_LLM=local 세션만 40_IP/** 와 confidential 노트를 읽고 쓴다. Read/Grep/Glob Guard 추가, MCP 는 로컬 프로세스가 아니면 expose_confidential 을 무시", "ADR-0011 · Rev 2.1 §4.3"],
      ["Vault git 훅", "없음", "myelin init 이 pre-commit 훅(myelin doctor --staged)을 설치. AI 초안이 허용 범위 밖이면 커밋 거부(ai-scope). 기존 Vault 는 myelin doctor --install-hook", "Rev 2.1 §3.1"],
      ["Commons 제안", "로컬 브랜치 커밋까지", "\"push 하고 PR 을 만든다\" 체크 → push + PR(gh). 실패해도 커밋 보존 + 다음 명령 안내. 인용 집계(top_contributors). 두 정체성(실명 기여 / 익명 지표) 문구", "ADR-0010 · ADR-0012"],
      ["도우미", "없음", "모든 화면 우측 하단 말풍선 → 도움말 챗봇. 로컬 Ollama 가 있으면 LLM 모드, 없으면 지식 기반 모드. 읽기 전용", "ADR-0013"],
      ["Obsidian", "테마 + cssclasses", "+ Graph View 설정(core-plugins · graph · app.json)과 Domain 별 Emboss 색 그룹, myelin obsidian setup|status, Dashboard 의 Obsidian Graph View 카드, Vault 등록", "v1.3.1"],
      ["온보딩", "Pack 에서 Domain 선택", "Pack 에서 고르거나 직접 정의(개인 Pack me.* 가 자동 생성). Domain Packs 의 \"개인 Pack 으로 복사\"", "v1.1.1"],
      ["설치", "plugin/MCP 자동 등록", "--claude 를 줄 때만 plugin/MCP 등록(opt-in). install_obsidian.sh. 새 Vault 에 테마·Graph View 설정 선설치", "v1.1.0 · v1.3.1"],
      ["Guard", "Write/Edit + Bash", "+ Read/Grep/Glob(기밀 라우팅). Bash 우회 경로 확대(stdin 스크립트 · find -delete · curl -o · 리다이렉션 …). Write/Edit 는 .md 만", "v1.1.1 · v1.2.0"],
      ["Hub 포트", "예시 8801", "기본 8787(환경변수 MYELIN_HUB_PORT)", "—"],
  ],
  [2.8, 3.6, 7.6, 3.0])

T("menus", "1 메뉴 개요",
  ["메뉴", "하는 일", "쓰기 여부", "처음 사용자가 먼저 볼 것"],
  [
      ["온보딩 마법사", "Vault 를 처음 만든다 (이름 · 3+1 Domain · Level). Domain 은 Pack 에서 고르거나 직접 정의", "Vault 생성 1회 ([myelin:init] 커밋)", "이름과 Primary Domain 만 정하면 된다"],
      ["Dashboard", "KPI 타일 · Myelination 강도 분포 · 지표 9종과 목표 판정 · Obsidian Graph View 카드 · 최근 커밋", "Graph View 설정 쓰기 버튼만 (commit 1개)", "verified 숫자와 Review Gate 대기 수"],
      ["My Domains", "Domain Orb → Domain 상세 · Sub-MOC 주제 검색 · 노트 표(강도) · Obsidian 에서 MOC 열기", "없음 (읽기 전용)", "Primary Domain 의 Orb"],
      ["Inbox / Review Gate", "AI 초안을 읽고 Promote/Reject, 패치를 미리보기 후 Apply, Commons 에 제안(push·PR 선택)", "클릭마다 git commit 1개", "\"내용 읽기\" 다음의 Promote 버튼"],
      ["Domain Packs", "설치된 Pack 을 보고, Reference Pack 을 개인 Pack 으로 복사하거나 개인 Pack(me.*) 을 편집·저장", "개인 Pack 저장 시 commit", "Reference Pack 의 Domain 구성"],
      ["Settings", "Level · MCP 스위치 · Commons(로컬 clone · 원격 · 연동 · 익명 지표 opt-in) · Skill 개방 표", "저장 시 commit", "Level 1 로 시작"],
      ["도우미 (우측 하단)", "MYELIN 개념·절차와 이 Vault 의 상태를 묻고 답한다. 근거 노트를 Obsidian 링크로 준다", "없음 (읽기 전용, 대화 미저장)", "\"AI 초안을 승격하는 절차는?\" 같은 예시 질문"],
  ],
  [3.0, 7.2, 3.4, 3.4])

T("domains", "2 3+1 Domain 규칙",
  ["역할", "뜻", "예 (Reference Pack · 직접 정의)", "선택"],
  [
      ["Primary", "지금 내 업무의 근간. 매주 기록할 거리가 생기는 Domain", "D1 AI/ML/AX Platform · CLOUD-SEC(직접 정의)", "필수"],
      ["Research", "스스로 파고드는 연구 주제. 논문·실험·사이드 프로젝트", "D3 자율학습 Robotics", "권장"],
      ["Growth", "1~2년 안에 역량을 갖추고 싶은 영역", "D4 Auto Design (SolidWorks)", "권장"],
      ["+1 (extra)", "3개월 운영 후 지표로 근거가 생기면 add-domain 으로 추가", "D2 Embedded SW · D4 AI_Project(개인 Pack)", "선택"],
  ],
  [2.6, 6.4, 5.4, 2.6])

T("kpi", "3 Dashboard KPI",
  ["타일·지표", "뜻", "목표 (3개월)", "읽는 법"],
  [
      ["verified · myelinated", "사람이 승격한 노트 수", "주당 +5 (구성원) / +10 (Brian)", "늘어나면 지식 경로가 굵어진 것"],
      ["Review Gate 대기", "승격을 기다리는 AI 초안·패치 수", "0 에 가깝게", "쌓이면 Review Gate 로 이동"],
      ["Guard 차단 (7일)", "Claude 가 사람 승인 영역에 쓰려다 막힌 횟수", "참고용", "많으면 Skill 사용 방식 점검"],
      ["doctor 파싱 오류", "frontmatter 가 깨진 노트 수", "0", "myelin doctor 로 원인 확인"],
      ["inbox_dwell_days", "Inbox 노트의 평균 체류일", "≤ 7", "/distill 로 분해하거나 Archive"],
      ["verified_per_week", "최근 7일 승격 수", "≥ 5 (구성원) / ≥ 10 (Brian)", "매주 Review Gate 를 비운다"],
      ["ai_promotion_rate", "AI 초안 중 승격된 비율", "0.40 ~ 0.70", "0.70 초과면 검토 부실 신호 → 목표 이탈"],
      ["orphan_ratio", "링크 없는 노트 비율", "≤ 0.10", "/link 제안으로 연결"],
      ["cross_domain_ratio", "Domain 2개 이상 노트 비율", "≥ 0.15", "/compose · /ip-scan 의 재료"],
      ["outputs_per_month", "80_Outputs 산출물 수 (30일)", "≥ 2", "Vault 가 실제 일에 재투입되는지"],
      ["reuse_ratio (v2.0)", "verified 노트 중 강도 ≥ 2 인 비율", "≥ 0.30", "verified 가 재사용되고 있는지 — 복리의 지표"],
      ["reuse_lag_days (v2.0)", "재료 노트 created → 그 노트를 인용한 산출물 created 의 중앙값(일)", "≤ 30", "길면 승격한 지식이 일에 늦게 쓰이는 것"],
      ["compose_sources_avg (v2.0)", "산출물 1건당 verified 재료 노트 수의 평균", "≥ 3", "낮으면 /compose 가 Vault 를 재료로 쓰지 않는 것"],
  ],
  [3.6, 5.6, 3.6, 4.2])

T("strength", "3b 미엘린 강도",
  ["강도", "라벨", "뜻", "조건 (ADR-0012)"],
  [
      ["0", "(표시 없음)", "verified 가 아닌 노트", "draft · unverified · deprecated"],
      ["1", "myelin ×1", "verified — 검증만 됨", "Review Gate 에서 승격"],
      ["2", "myelin ×2", "reused — 재사용됨", "다른 노트 2곳 이상에서 링크되거나, /compose 산출물의 재료(source)로 1회 이상 쓰임"],
      ["3", "myelin ×3", "hub — 신경망의 허브", "Commons 에서 1회 이상 인용되었거나, 링크 5곳 이상이고 산출물 재료 2회 이상"],
  ],
  [1.4, 2.6, 4.4, 8.6])

T("chips", "4 상태 칩 (색 의미)",
  ["칩·표시", "색", "뜻", "누가 만드나"],
  [
      ["draft", "Violet 외곽선", "사람이 쓰는 중인 초안", "사람"],
      ["ai-draft", "Cyan 채움 + 카드 glow", "Claude 가 만든 초안 (origin: claude). 항상 Cyan 으로 구분", "Claude (Skill/MCP)"],
      ["verified", "Lime 채움 + Lime 링", "사람이 Review Gate 에서 승격한 노트 (myelinated)", "사람의 클릭"],
      ["unverified", "회색 외곽선", "사람이 썼지만 아직 검토 전", "사람"],
      ["blocked / invalid", "Magenta 외곽선 배지", "Guard 차단, doctor 위반, 깨진 frontmatter, Commons 공개 불가", "시스템"],
      ["myelin ×N (v2.0)", "×1 회색 · ×2 ×3 Lime 채움 + 글자", "노트의 미엘린 강도 1~3 (3b 표)", "시스템 (색인이 계산)"],
      ["달성 / 목표 이탈 (v2.0)", "Lime 채움 / Magenta 외곽선", "Dashboard 지표가 3개월 목표 안인지 밖인지", "시스템"],
      ["publishable / blocked (Commons)", "Lime / Magenta", "Commons 후보가 하드 게이트와 비식별 검사를 통과했는지", "시스템"],
      ["Domain Orb accent", "Pack 의 accent 색", "Domain 정체성 (violet/lime/cyan/magenta) — Obsidian Graph View 색 그룹과 동일", "Pack"],
  ],
  [3.6, 3.8, 6.4, 3.2])

T("review", "5 Review Gate 동작",
  ["버튼", "조건", "결과", "git commit"],
  [
      ["내용 읽기 (v2.0)", "ai-drafts 의 초안", "본문을 연다. 이 뒤에만 Promote 가 켜진다 (읽음 확인 장치)", "없음"],
      ["Promote", "본문을 읽었고, frontmatter 가 유효하며, Domain 이 설정됨. 같은 이름 노트가 Domain 폴더에 없음(있으면 409)", "10_Domains/<D>/ 로 이동, status verified, reviewed 오늘, origin claude→mixed. myelinated 애니메이션", "[myelin:review] promote …"],
      ["Reject", "ai-drafts 의 초안", "99_Archive/rejected/ 로 이동, status deprecated", "[myelin:review] reject …"],
      ["미리보기 (패치, v2.0)", "_patches/ 의 *.patch.md", "diff · fail-closed 검사 결과(40_IP · confidential · verified→미검증 AI 링크 · 결과 frontmatter) · digest 표시", "없음"],
      ["Apply (패치)", "미리보기의 digest 와 일치 (불일치 409)", "대상 노트에 append(under 구획) / replace / frontmatter set 적용, 패치는 _patches/applied/ 로 이동. 결과가 frontmatter 를 깨뜨리면 되돌림", "[myelin:link] apply patch …"],
      ["Reject (패치)", "proposed 패치", "_patches/rejected/ 로 이동", "[myelin:review] reject patch …"],
      ["Commons 제안", "share:true + verified + ≤ internal + 40_IP 밖 + 기밀 링크 없음 + 비식별 통과 + \"고객명·IP·기밀이 없음을 확인했습니다\" 체크. Settings › Commons 의 로컬 clone 필요", "Commons clone 의 publish/<id> 브랜치에 커밋. \"push 하고 PR 을 만든다\" 를 체크하면 push + PR(gh)까지. 실패해도 커밋은 남고 다음 명령을 안내", "[myelin:publish] …"],
  ],
  [2.8, 5.0, 6.0, 3.2])

T("levels", "6 Level 과 Skill",
  ["Level", "이름", "열리는 Skill", "AI 가 쓸 수 있는 경로"],
  [
      ["0", "read-only", "없음 (읽기·검색만)", "없음"],
      ["1", "drafts", "/ingest /distill /interview /review", "00_Inbox/ai-drafts, 50_Journal"],
      ["2", "link-patch", "+ /link /compose", "+ _patches, 80_Outputs/drafts"],
      ["3", "scheduled review", "+ /ip-scan", "+ 40_IP/candidates (MYELIN_LLM=local 세션에서만 — 외부 세션은 Level 3 이어도 차단)"],
  ],
  [1.4, 2.8, 5.4, 7.4])

T("skills", "7 Skill 명령",
  ["Skill", "입력", "만드는 것", "저장 위치", "Level", "v2.0 변경"],
  [
      ["/myelin:ingest", "파일·URL·PDF", "literature 노트 초안 (출처·핵심 주장·인용 위치)", "00_Inbox/ai-drafts", "1", "—"],
      ["/myelin:distill", "Inbox 노트 또는 80_Outputs 산출물", "원자(evergreen) 노트 초안, 중복 검사, 용어 정규화. 산출물을 주면 ## 환류 후보 절을 lesson/evergreen/adr 초안으로 되돌림 (tags: [refeed])", "00_Inbox/ai-drafts", "1", "환류 입력 추가"],
      ["/myelin:interview", "Pack 질문 은행", "인터뷰 → lesson 초안", "00_Inbox/ai-drafts", "1", "—"],
      ["/myelin:review", "doctor·stats", "주간 품질 보고서 — 지표 9종 이탈, 환류 대기 산출물, doctor 할 일", "50_Journal", "1", "환류 대기·재사용 지표 보고"],
      ["/myelin:link", "노트·Domain", "링크 후보 + MOC 갱신 패치 제안 (*.patch.md)", "_patches", "2", "under 구획 · set 허용 키"],
      ["/myelin:compose", "목적·독자", "산출물 초안 (설계서·제안서·보고서) + ## 환류 후보 절", "80_Outputs/drafts", "2", "환류 후보 절"],
      ["/myelin:ip-scan", "cross-domain verified 노트", "특허 후보 초안", "40_IP/candidates", "3", "로컬 세션 전용 (ADR-0011)"],
  ],
  [2.8, 3.0, 5.2, 2.6, 1.0, 2.4])

T("cli", "8 CLI 명령",
  ["명령", "하는 일", "비고"],
  [
      ["myelin init --vault <경로>", "Vault 생성 (3+1 마법사). git init + [myelin:init] 커밋 + pre-commit 훅 + Obsidian 테마·Graph View 설정", "기존 파일 있으면 중단 · --no-git · --dry-run · -y 비대화형"],
      ["myelin add-domain <ID>", "Pack 의 Domain 을 +1 로 추가. Graph View 색 그룹도 추가", "기존 노트 불변 · [myelin:add-domain] 커밋 · --pack 으로 중복 ID 선택"],
      ["myelin doctor", "무결성 점검 규칙 8종 + ai-scope", "오류 시 exit 1 · --json · --strict · --staged(훅용) · --install-hook"],
      ["myelin stats", "개인 지표 9종 + 강도 분포", "--json · --today YYYY-MM-DD"],
      ["myelin skills / interview-bank", "Level 별 Skill 표 / 질문 은행", "Skill 이 내부에서 호출 · --json"],
      ["myelin publish <노트>", "Commons 제안 (비식별 검사 → 확인 → 브랜치 커밋)", "--check 검사만 · --push · --pr (gh) · -y"],
      ["myelin pack-propose <packs/me.x>", "개인 Pack 을 Org Pack 후보로 제안 (propose/<name> 브랜치)", "--commons · -y"],
      ["myelin metrics show|send|aggregate", "익명 지표 (opt-in). aggregate 는 인용 집계(reuse · top_contributors) 포함", "숫자와 id 목록만 전송"],
      ["myelin commons init|add-pack|registry", "Commons 저장소·Registry", "운영자용"],
      ["myelin obsidian setup [--open] (v2.0)", ".obsidian/ 에 Graph View 설정을 쓴다. --open 은 obsidian:// 로 Vault 를 열어 등록", "있는 파일은 병합, 읽을 수 없는 JSON 이면 중단"],
      ["myelin obsidian status (v2.0)", "Graph View 설정 · Domain 색 그룹 · 테마 · Obsidian 등록 여부", "Hub Dashboard 카드와 같은 정보"],
      ["myelin-hub --vault <경로> [--port 8787]", "Hub 실행", "localhost 전용 · MYELIN_HUB_PORT"],
      ["myelin-mcp --vault <경로>", "MCP 서버 (stdio)", "Claude Desktop/Code 에 등록 · MYELIN_VAULT"],
  ],
  [5.4, 6.4, 5.2])

T("folders", "9 Vault 폴더",
  ["폴더·파일", "들어가는 것", "AI 쓰기"],
  [
      ["00_Inbox/", "가공 전 메모, clips/", "ai-drafts/ 만"],
      ["00_Inbox/ai-drafts/", "Claude 초안 (origin claude, status draft)", "허용 (L1)"],
      ["10_Domains/<ID>/", "verified 노트, _MOC.md", "금지 (사람 승인 영역)"],
      ["20_Projects/ 30_Resources/", "프로젝트·문헌 노트", "금지"],
      ["40_IP/", "특허 후보 (confidential). MYELIN_LLM=local 세션만 읽고 쓴다", "candidates/ 만 (L3 + 로컬 세션)"],
      ["50_Journal/", "주간 리뷰", "허용 (L1)"],
      ["80_Outputs/", "산출물 (## 환류 후보 절)", "drafts/ 만 (L2)"],
      ["90_System/", "templates · prompts · taxonomy.md · deid.yaml · hooks/pre-commit", "금지"],
      ["99_Archive/", "보관 · rejected/ (거부된 초안)", "금지"],
      ["_patches/", "AI 의 변경 제안 *.patch.md · applied/ · rejected/ 이력", "허용 (L2, 직속 *.patch.md 만)"],
      [".obsidian/ (v2.0)", "themes/MYELIN Emboss · appearance.json · core-plugins.json · graph.json · app.json", "시스템 (init / obsidian setup)"],
      [".myelin/", "색인(index.db) · 로그 · 지표 이벤트", "시스템"],
      ["myelin.yaml · CLAUDE.md · .claude/settings.json", "설정 · Claude 운영 규칙 · 2차 방어 deny 규칙", "금지 (사람만)"],
  ],
  [5.0, 7.4, 4.6])

T("trouble", "10 문제 해결",
  ["증상", "원인", "해결"],
  [
      ["Hub 가 403 을 돌려준다", "localhost 가 아닌 주소, 또는 다른 사이트 탭에서 호출 (Host/Origin 검사)", "http://127.0.0.1:<포트> 로 접속"],
      ["Promote 버튼이 비활성", "아직 \"내용 읽기\" 를 누르지 않았거나, frontmatter 가 깨진 초안 (invalid 배지)", "내용 읽기 → Promote. invalid 면 Obsidian 에서 frontmatter 수정 후 새로고침"],
      ["Promote 가 409", "같은 이름의 노트가 Domain 폴더에 이미 있다", "초안 제목을 바꾸거나 기존 노트를 보강하는 패치로 전환"],
      ["Apply 가 409", "미리보기 뒤 패치 파일이 바뀌어 digest 가 다르다", "미리보기를 다시 열어 검사 결과를 확인하고 Apply"],
      ["Guard 가 전부 막는다", "myelin.yaml 없음 또는 level 0", "Vault 폴더에서 실행, Settings 에서 Level 1 이상"],
      ["40_IP · confidential 노트를 Claude 가 읽지 못한다", "외부 세션 (MYELIN_LLM=local 없음) — 의도된 동작", "로컬 LLM 세션에서만 MYELIN_LLM=local 을 설정. 외부 Claude 에는 보내지 않는다"],
      ["git commit 이 거부된다 (pre-commit)", "origin: claude 노트가 허용 범위 밖에 있거나, reviewed 없는 verified, confidential+share", "메시지의 규칙 이름으로 노트를 고친다. --no-verify 는 Guard 가 막는다"],
      ["Commons 제안이 blocked", "share:false / 미검증 / confidential / 비식별 적중", "원인 줄을 읽고 노트를 고친다. 별명·프로젝트 코드는 90_System/deid.yaml 에 등록"],
      ["Commons 연동을 켤 수 없다", "commons.remote 가 비어 있다", "Settings › Commons 에 원격 저장소를 먼저 적는다. 로컬 clone 은 myelin commons init <path>"],
      ["도우미가 \"지식 기반 모드\"", "로컬 Ollama(http://127.0.0.1:11434)가 꺼져 있거나 모델이 없다", "ollama serve 후 다시 연다. MYELIN_ASSISTANT_URL · MYELIN_ASSISTANT_MODEL 로 변경"],
      ["Hub 의 Obsidian 링크가 열리지 않는다", "Obsidian 이 이 Vault 를 모른다 (미등록)", "Dashboard 카드의 \"Obsidian 에서 Vault 열기\" 를 한 번 누르거나 myelin obsidian setup --open. Obsidian 에서 Open folder as vault 도 된다"],
      ["Graph View 에 색이 없다", ".obsidian/graph.json 에 Domain 색 그룹이 없다", "Dashboard › Graph View 설정 쓰기 또는 myelin obsidian setup"],
      ["검색이 비어 있다", "2자 이하 토큰은 부분 일치, 3자 이상은 FTS", "단어를 공백으로 나눠 입력"],
      ["doctor 오류", "규칙 8종(schema · fail-closed · orphan · confidential-share · broken-link · duplicate-id · domain · temporal) + ai-scope 중 하나 위반", "메시지의 규칙 이름으로 노트 수정"],
  ],
  [4.4, 6.0, 6.6])

T("install", "11 설치 옵션",
  ["플래그", "뜻", "기본"],
  [
      ["--vault <path>", "Obsidian 테마 복사 + myelin obsidian setup (Graph View 설정). 기존 파일은 .bak 백업 후 갱신", "없음 (생략 시 테마 단계 생략)"],
      ["--claude", "Claude Code plugin(로컬 marketplace myelin-local)과 MCP 서버 등록. 이 플래그가 없으면 등록하지 않는다 (opt-in)", "꺼짐"],
      ["--bin <dir>", "myelin · myelin-mcp · myelin-hub 런처를 둘 폴더", "~/.local/bin"],
      ["--no-web", "Hub 웹 빌드 생략 (npm 없으면 자동 생략)", "빌드함"],
      ["--no-plugin / --no-mcp", "--claude 일 때 plugin / MCP 등록만 각각 생략", "—"],
      ["--dry-run", "실행할 명령만 출력", "꺼짐"],
      ["scripts/install_obsidian.sh --vault <path>", "Obsidian 자체를 설치(.deb, Debian/Ubuntu, sudo)하고 테마를 넣은 뒤 Vault 를 연다", "별도 스크립트"],
  ],
  [5.0, 8.6, 3.4])

T("assistant", "12 도우미",
  ["항목", "내용"],
  [
      ["여는 곳", "Hub 모든 화면의 우측 하단 \"도우미\" 말풍선 → 모달. 닫기 또는 ESC"],
      ["답하는 것", "MYELIN 의 개념·사용법·절차·명령, 그리고 이 Vault 의 상태 — Level · Domain · KPI 이탈 · 검토 대기 · 관련 노트. 근거 노트는 Obsidian 링크(obsidian://)로 준다"],
      ["하지 않는 것", "노트 쓰기 · 승격 · 패치 적용 · 공개 · 커밋 · 삭제. 안내만 하며 행위는 Hub 버튼과 Claude Code Skill 로 사람이 한다. 대화는 서버에 저장하지 않는다"],
      ["LLM 모드", "로컬 Ollama(기본 http://127.0.0.1:11434)가 켜져 있으면 LLM 이 문장을 만든다. 모델 기본값은 exaone3.5 → qwen2.5 … 순으로 있는 것을 쓴다"],
      ["지식 기반 모드", "Ollama 가 없으면 내장 지식 절(assistant_kb.md 17절)을 질문 키워드로 골라 그대로 보여 준다. 헤더에 \"지식 기반 모드\" 칩"],
      ["기밀 규칙", "엔드포인트가 loopback(127.0.0.1)이 아니면 confidential 노트와 40_IP/** 는 컨텍스트에서 제외한다 (ADR-0011 과 같은 판정)"],
      ["환경변수", "MYELIN_ASSISTANT_URL (Ollama 주소) · MYELIN_ASSISTANT_MODEL (모델 이름)"],
      ["입력 가드", "프롬프트 인젝션 · 비밀값(토큰·키) 요청은 거부. 한글 IME 조합 중 Enter 는 무시"],
      ["예시 질문", "\"AI 초안을 승격(Promote)하는 절차는?\" · \"Level 2 에서 쓸 수 있는 Skill 은?\" · \"기밀 노트는 외부 LLM 에 가나요?\" · \"Commons 에 노트를 공유하는 방법은?\""],
  ],
  [3.2, 13.8])

T("obsidian", "13 Obsidian 설정 파일",
  ["파일 (<Vault>/.obsidian/)", "내용", "만드는 때"],
  [
      ["themes/MYELIN Emboss/", "manifest.json · theme.css — design/tokens.json 에서 생성된 Emboss 테마", "myelin init · Hub 온보딩 · install.sh --vault"],
      ["appearance.json", "라이트 모드 + MYELIN Emboss 선선택 (cssTheme)", "같음"],
      ["core-plugins.json (v2.0)", "Graph · Backlinks · Outgoing links · Tags · Properties · Templates 켜짐. Daily notes · Sync · Publish 꺼짐", "myelin init · myelin obsidian setup · Dashboard 버튼"],
      ["graph.json (v2.0)", "Domain 폴더(10_Domains/<ID>/)마다 그 Domain 의 Emboss accent 색 그룹 + 00_Inbox(muted) · 40_IP(magenta) · 99_Archive(border). 태그 표시, 미해결 링크 숨김", "같음. add-domain 이 새 Domain 의 그룹을 보탠다"],
      ["app.json (v2.0)", "새 노트 → 00_Inbox, [[stem]] 위키링크, 이름 변경 시 링크 자동 갱신", "같음"],
  ],
  [4.6, 8.0, 4.4])

T("cssclasses", "14 cssclasses",
  ["cssclasses 값", "제목에 그려지는 것", "쓰는 때"],
  [
      ["mx-status-verified", "라임 링 + verified 칩 (myelinated)", "Review Gate 에서 승격한 노트. 승격 시 자동으로 붙지는 않으므로 필요할 때 적는다"],
      ["mx-status-ai-draft", "시안 glow + ai-draft 칩", "00_Inbox/ai-drafts 의 초안을 Obsidian 에서 구분해 보고 싶을 때"],
      ["mx-status-blocked", "마젠타 외곽선 + blocked 칩", "doctor 가 지적했거나 Guard 에 막힌 노트를 표시해 둘 때"],
  ],
  [4.0, 5.0, 8.0])

T("guard", "15 Guard (Claude Code)",
  ["Hook", "대상 도구", "검사 내용", "차단 시"],
  [
      ["guard_write.py", "Write · Edit · MultiEdit · NotebookEdit", "Level 쓰기 범위(6 표) 밖 경로, Vault 밖, 심볼릭 링크, .., 숨김 디렉터리, .md 가 아닌 파일, _patches/applied|rejected/. 외부 세션의 40_IP 쓰기", "exit 2 — 도구 호출 거부, 사유 표시"],
      ["guard_bash.py", "Bash", "리다이렉션·tee·sed -i·find -delete·stdin 스크립트(bash -s, curl | sh)·Vault 내부 스크립트 실행·wget/curl -o 등 쓰기 우회. git 은 읽기 + add/commit 만 (--no-verify 차단). grep/find/git show 로 기밀 읽기", "같음"],
      ["guard_read.py (v2.0)", "Read · Grep · Glob", "MYELIN_LLM=local 이 아닌 세션에서 40_IP/** 와 classification: confidential 노트 접근", "같음 — 로컬 세션에서만 허용"],
      [".claude/settings.json", "(2차 방어)", "Write/Edit deny 목록 + Read(./40_IP/**) deny. bash/sh/perl/ruby/curl/wget 은 ask", "Claude Code 가 확인을 묻거나 거부"],
  ],
  [2.8, 3.2, 8.0, 3.0])

T("mcp", "16 MCP 도구",
  ["도구", "설명", "권한"],
  [
      ["search_notes(query, domain?, status?, type?, limit)", "FTS5(+LIKE) 검색, 스니펫·순위", "read"],
      ["get_note(ref)", "경로 또는 id → frontmatter + 본문 + backlinks", "read (confidential 은 본문 비공개)"],
      ["query_frontmatter(type?, status?, origin?, classification?, domain?, tag?, share?, created/reviewed 범위, order, limit)", "Dataview 류 질의", "read"],
      ["list_domains()", "3+1 Domain · 역할 · 노트 수 · Pack 메타", "read"],
      ["vault_stats(today?)", "지표 9종 + 강도 분포", "read"],
      ["write_draft(title, body, type, domain, tags, source)", "00_Inbox/ai-drafts/ 에 새 초안", "opt-in: mcp.write_draft: true + level ≥ 1"],
  ],
  [7.0, 6.0, 4.0])

# ----------------------------------------------------------------------------- 그림
FIGURES = {
    "dashboard": ("dashboard.png", P6 / "01-dashboard.png"),
    "onboarding": ("onboarding.png", P6 / "00-onboarding.png"),
    "domains": ("domains.png", P6 / "02-my-domains.png"),
    "review": ("review.png", P6 / "05-review-gate.png"),
    "promoted": ("promoted.png", P6 / "06-promoted.png"),
    "packs": ("packs.png", P6 / "03-domain-packs.png"),
    "settings": ("settings.png", P6 / "04-settings.png"),
    "assistant": ("assistant.png", P6 / "assistant.png"),
    "obs-verified": ("obsidian-verified.png", V1_IMG / "obsidian-03-verified.png"),
    "obs-ai-draft": ("obsidian-ai-draft.png", V1_IMG / "obsidian-04-ai-draft.png"),
    "emboss-top": ("emboss-gallery-top.png", V1_IMG / "emboss-gallery-top.png"),
    "emboss-bottom": ("emboss-gallery-bottom.png", V1_IMG / "emboss-gallery-bottom.png"),
}


# ----------------------------------------------------------------------------- docx helpers
def set_run_font(run, size=None, bold=None, color=None, mono=False, italic=None):
    name = MONO if mono else FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_borders(cell, color=BORDER, sz="4"):
    tcpr = cell._element.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcpr.append(borders)


def para_shading(p, fill):
    ppr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)



def fix_table_widths(t, widths):
    """LibreOffice/Word 가 열 너비를 지키도록 tblLayout fixed + tblGrid + 열 너비를 함께 쓴다."""
    tbl = t._tbl
    tblpr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    total = int(sum(widths) * 567)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        grid.append(gc)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)


class Doc:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, m, Cm(2.0))
        normal = self.d.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.font.color.rgb = RGBColor.from_string(INK)
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        for lvl, size, color in ((1, 18, VIOLET), (2, 14, INK), (3, 11.5, INK)):
            st = self.d.styles[f"Heading {lvl}"]
            st.font.name = FONT
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor.from_string(color)
            r = st.element.get_or_add_rPr()
            rf = r.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                r.append(rf)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(attr), FONT)
            st.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        self.fig_no = 0
        self.tab_no = 0
        self._footer()

    def _footer(self):
        s = self.d.sections[0]
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"MYELIN 사용자 가이드 {GUIDE_VER} · {GUIDE_DATE} · {ORG} · ")
        set_run_font(r, 8, color=MUTED)
        # page number field
        r2 = p.add_run()
        set_run_font(r2, 8, color=MUTED)
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
                r2._element.append(el)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                r2._element.append(it)

    # --- blocks
    def h(self, level, text):
        p = self.d.add_heading(text, level=level)
        for r in p.runs:
            set_run_font(r, bold=True, color=VIOLET if level == 1 else INK)
        return p

    def p(self, text="", size=10, color=INK, bold=None, align=None, italic=None, after=None):
        p = self.d.add_paragraph()
        if align:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        self._rich(p, text, size, color, bold, italic)
        return p

    def _rich(self, p, text, size=10, color=INK, bold=None, italic=None):
        """`code` 조각은 Consolas 로, **굵게** 는 bold 로."""
        import re
        parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1])
                set_run_font(r, size - 0.5, bold=bold, color=color, mono=True)
            elif part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                set_run_font(r, size, bold=True, color=color, italic=italic)
            else:
                r = p.add_run(part)
                set_run_font(r, size, bold=bold, color=color, italic=italic)

    def bullets(self, items, numbered=False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            p = self.d.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(3)
            if isinstance(it, tuple):
                lead, rest = it
                r = p.add_run(lead + " ")
                set_run_font(r, 10, bold=True, color=INK)
                self._rich(p, rest)
            else:
                self._rich(p, it)

    def code(self, text):
        for line in text.rstrip("\n").split("\n"):
            p = self.d.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Cm(0.4)
            para_shading(p, SURFACE_2)
            r = p.add_run(line if line else " ")
            set_run_font(r, 9, color=INK, mono=True)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def callout(self, title, text):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        shade(c, VIOLET_50)
        set_cell_borders(c, color=VIOLET, sz="6")
        p = c.paragraphs[0]
        r = p.add_run(title)
        set_run_font(r, 10, bold=True, color=VIOLET)
        p2 = c.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        self._rich(p2, text, 9.5)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, key, note=True):
        sheet, headers, rows, widths = TABLES[key]
        self.tab_no += 1
        cap = self.d.add_paragraph()
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.keep_with_next = True
        import re as _re
        r = cap.add_run(f"표 {self.tab_no}. {_re.sub(r'^\d+b? ', '', sheet)}")
        set_run_font(r, 9, bold=True, color=INK)
        if note:
            r2 = cap.add_run(f"  — 동반 파일 {XLSX_NAME} 시트 \"{sheet}\" 와 동일")
            set_run_font(r2, 8.5, color=MUTED)
        t = self.d.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        fix_table_widths(t, widths)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.width = Cm(widths[i])
            shade(c, SURFACE_2)
            set_cell_borders(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            rr = p.add_run(h)
            set_run_font(rr, 9, bold=True, color=INK)
        # repeat header row
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                c = cells[i]
                c.width = Cm(widths[i])
                set_cell_borders(c)
                if ri % 2 == 1:
                    shade(c, VIOLET_50)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                self._rich(p, val, 8.5, INK, bold=(i == 0 and key not in ("assistant",)) or None)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, key, caption, width_cm=17.0, max_h_cm=22.0):
        from PIL import Image
        name, src = FIGURES[key]
        path = IMG / name
        self.fig_no += 1
        with Image.open(path) as im:
            w, h = im.size
        wc = width_cm
        hc = wc * h / w
        if hc > max_h_cm:
            hc = max_h_cm
            wc = hc * w / h
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(wc), height=Cm(hc))
        cap = self.d.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"그림 {self.fig_no}. {caption}")
        set_run_font(r, 9, color=MUTED)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------- 본문
def build_docx() -> Path:
    D = Doc()
    d = D.d

    # 표지
    for _ in range(4):
        D.p("", after=0)
    D.p("MYELIN", 36, VIOLET, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    D.p("Personal Knowledge OS — powered by Obsidian + Claude", 13, INK, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    D.p("“쓸수록 빨라지는 나의 지식 신경망”", 12, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    D.p("사용자 가이드 — Hub 메뉴 순서로 읽는 처음 사용자용 안내", 14, INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    D.p(f"{GUIDE_VER} · {GUIDE_DATE} · 제품 {PRODUCT_VER} 기준 · {ORG}", 10, MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    D.figure("dashboard", "MYELIN Hub — Dashboard (v1.3.1: Myelination 카드 · 지표 9종 · Obsidian Graph View 카드)", max_h_cm=15.5)
    D.page_break()

    # 차례
    D.h(1, "차례")
    toc = [
        "0. 서두 — MYELIN 의 설계 철학 · v1.0 에서 달라진 것",
        "1. 시작하기 — 설치와 실행",
        "2. 온보딩 마법사 — Vault 만들기",
        "3. Dashboard — 지식 신경망의 현재",
        "4. My Domains — Domain Map 과 MOC",
        "5. Inbox / Review Gate — 승격은 읽은 사람의 클릭으로",
        "6. Domain Packs — Domain 은 코드가 아니라 선언",
        "7. Settings — Level · MCP · Commons",
        "8. 도우미 — 우측 하단의 도움말 챗봇",
        "9. 매일·매주 — 지식 경로를 굵게 만드는 습관",
        "10. Obsidian 과 Claude Code 에서",
        "11. 부록 — 폴더 · 명령 · 문제 해결 · 제약 · 더 읽을 문서",
    ]
    for t in toc:
        D.p(t, 10.5, INK, after=3)
    D.p(f"동반 파일: {XLSX_NAME} (이 문서의 표 {len(TABLES)}개) · img/*.png (그림 {len(FIGURES)}장)", 9, MUTED)
    D.page_break()

    # 0
    D.h(1, "0. 서두 — MYELIN 의 설계 철학")
    D.callout("MYELIN ; Personal Knowledge OS — powered by Obsidian + Claude",
              "“쓸수록 빨라지는 나의 지식 신경망” — 개인의 관심·경험을 재사용 가능한 지식으로 전환하는 지식 운영체계. "
              "“미엘린(MYELIN)은 반복해서 쓰는 신경 경로를 감싸 신호를 빠르게 만드는 조직입니다. 경험을 반복해 연결하고 재사용할수록 지식 경로가 빨라진다는 컨셉입니다.” — Brian, Research Director, A3-AI WG")
    D.p("미엘린(수초)은 신경세포의 축삭을 감싸 신호 전달 속도를 높이는 절연층이다. 반복해서 쓰는 신경 경로일수록 미엘린화가 진행되어 더 빠르고 정확해진다. MYELIN 은 이 메커니즘을 지식에 이식한다. 경험을 노트로 감싸고(캡처), 반복해서 연결·재사용할수록(Distill·Link·Compose) 그 지식 경로가 빨라진다. 노트가 검증을 거쳐 verified 가 되는 순간 Hub 는 라임 링을 그리고, 그 노트가 다른 노트와 산출물에서 다시 쓰일수록 **미엘린 강도**(×1 → ×2 → ×3)가 오른다. verified 는 시작이지 끝이 아니다.")
    D.p("브라이언이 MYELIN 에 담은 의도는 세 가지이며, v2.0 은 각 의도를 문장이 아니라 **장치**로 지키도록 보강된 상태를 다룬다.")
    D.bullets([
        ("축적이 복리로 작동하는 사람.", "매일의 메모가 원자 노트가 되고, 원자 노트가 MOC 로 묶이고, 묶인 지식이 설계서·제안서·특허로 재투입되는 순환. 재투입된 산출물의 새 통찰은 **환류**(`## 환류 후보` → `/distill`)로 다시 원자 노트가 된다. 재사용 비율·재사용 지연·산출물당 출처 지표가 이 순환이 실제로 도는지 잰다."),
        ("AI 는 초안을 만들고 사람이 승격한다.", "Claude 는 `00_Inbox/ai-drafts` 에만 쓸 수 있고, 사람 승인 영역은 Guard 가 차단한다. 승격은 Review Gate 에서 **본문을 읽은 사람**의 클릭으로만 일어나며 git commit 으로 기록된다. Claude Code 밖에서 온 변경도 Vault 의 pre-commit 훅이 같은 규칙으로 검사한다."),
        ("Vault 는 개인의 것이고, 공유는 선택이다.", "조직은 개인 Vault 를 열람하지 않는다. 기밀(`confidential`, `40_IP`)은 `MYELIN_LLM=local` 세션에서만 읽히고 외부 LLM 에는 가지 않는다. 공유는 `share: true` 로 표시하고 비식별 검사를 통과한 노트에 한하며, Commons 에서는 작성자와 원 id 가 실명으로 남아 기여가 평판이 된다. 반면 지표는 익명이다(두 정체성)."),
    ])
    D.p("이 가이드는 Hub(웹 대시보드)의 메뉴 순서를 따른다. 처음 쓰는 사람은 온보딩 마법사부터 Review Gate 의 첫 승격까지를 15분 안에 마치는 것을 목표로 한다. 명령줄(CLI)과 Claude Code 의 Skill 은 각 메뉴에서 필요한 만큼만 언급하고 전체 목록은 부록에 둔다. 읽다가 막히면 Hub 우측 하단의 **도우미**에게 물으면 된다(8장).")
    D.table("menus")
    D.h(2, "0.1 v1.0 가이드에서 달라진 것")
    D.p("이 가이드의 v1.0(2026-10-18)은 제품 v1.0.0 을 기준으로 썼다. v2.0 은 제품 v1.1.0 ~ v1.3.1 에서 더해진 장치를 반영한다. 이미 v1.0 을 읽은 사람은 아래 표의 v2.0 열만 보면 된다.")
    D.table("changes")

    # 1
    D.h(1, "1. 시작하기 — 설치와 실행")
    D.h(2, "1.1 준비물")
    D.bullets([
        "Python 3.12, uv, git, Node 22 (Hub 웹 빌드), Google Chrome. Claude Code 는 Skill 과 MCP 를 쓸 때 필요하다.",
        "Obsidian (Vault 를 열고 노트를 쓰는 편집기). MYELIN 은 Obsidian 의 Markdown 파일을 정본으로 삼는다. 없으면 `scripts/install_obsidian.sh` 가 설치한다(Debian/Ubuntu).",
        "(선택) 로컬 Ollama — Hub 도우미를 LLM 모드로 쓰고 싶을 때. 없어도 도우미는 지식 기반 모드로 동작한다.",
    ])
    D.h(2, "1.2 설치")
    D.bullets([
        ("저장소를 받는다.", "`git clone git@github.com:wdlab1958/MYELIN.git myelin && cd myelin`"),
        ("설치 스크립트를 실행한다.", "`scripts/install.sh --vault ~/Vaults/me --claude` — Python/uv 를 검사하고, 패키지를 동기화하고, Hub 웹을 빌드하고, `~/.local/bin` 에 `myelin` · `myelin-mcp` · `myelin-hub` 런처를 만들고, Obsidian 테마와 Graph View 설정을 Vault 에 넣는다. **Claude Code plugin 과 MCP 는 `--claude` 를 줄 때만 등록한다**(opt-in). 기존 파일은 덮어쓰지 않고, 테마는 `.bak` 으로 백업한 뒤 갱신한다."),
        ("Hub 를 띄운다.", "`myelin-hub --vault ~/Vaults/me` 를 실행한 뒤 브라우저에서 `http://127.0.0.1:8787` 을 연다. 포트가 겹치면 `--port` 또는 환경변수 `MYELIN_HUB_PORT` 로 바꾼다. Hub 는 이 컴퓨터(localhost)에서만 열린다."),
    ], numbered=True)
    D.code("git clone git@github.com:wdlab1958/MYELIN.git myelin && cd myelin\nscripts/install.sh --vault ~/Vaults/me --claude\nmyelin-hub --vault ~/Vaults/me          # → http://127.0.0.1:8787")
    D.table("install")
    D.p("개발 중에는 FE(Vite)와 BE(API)를 따로 띄울 수 있다. FE 는 `/api` 요청을 BE 로 전달하므로 두 포트만 비어 있으면 된다.", 9.5, MUTED)
    D.h(2, "1.3 업그레이드")
    D.p("`git pull && scripts/install.sh --vault <Vault> [--claude]` 뒤 `myelin doctor`. 스크립트는 멱등이다. v1.2.0 이전에 만든 Vault 는 `myelin doctor --install-hook` 으로 pre-commit 훅을, v1.3.1 이전 Vault 는 `myelin obsidian setup --vault <Vault>` 로 Graph View 설정을 보탠다(설치 스크립트가 `--vault` 를 받으면 둘 다 자동).")

    # 2
    D.h(1, "2. 온보딩 마법사 — Vault 만들기")
    D.p("Hub 가 가리키는 폴더에 아직 Vault 가 없으면 온보딩 화면이 먼저 열린다. 이름과 Primary Domain 만 정하면 Vault 가 생기고, 첫 AI 초안 “MYELIN 시작하기” 가 Review Gate 에 올라온다. Vault 는 **빈 폴더**에서 새로 만들며, 다른 사람의 노트나 Domain 을 가져오지 않는다.")
    D.figure("onboarding", "온보딩 마법사 — 이름 · Domain 정하기(Pack 또는 직접 정의) · 3+1 역할 배정 · Level", max_h_cm=20)
    D.h(2, "2.1 3+1 Domain 규칙")
    D.p("Domain 은 지식을 묶는 큰 서랍이다. 시작부터 많이 만들면 MOC 가 비고 지속성이 떨어진다. 그래서 MYELIN 은 세 개로 시작하고 하나를 나중에 더하는 규칙을 둔다.")
    D.table("domains")
    D.h(2, "2.2 입력 항목")
    D.bullets([
        ("이름(필수) · 팀 · 이메일.", "이메일은 Commons 익명 지표의 owner 해시 재료로만 쓰이고 그대로 노출되지 않는다."),
        ("내 Domain 정하기.", "**조직 업무 Pack(A3)** 에서 고르거나, **직접 정의**한다 — ID(영문 대문자·숫자·하이픈, 예 `CLOUD-SEC`), 이름, Sub-MOC 주제(쉼표 구분), 한 줄 설명. 직접 정의한 Domain 은 개인 Pack `me.my-domains` 로 Vault 의 `packs/` 에 저장된다. Reference Pack(`brian.physical-ai`)은 한 엔지니어의 지식 지형도를 참고용으로 보여 줄 뿐 따라 할 필요가 없다."),
        ("3+1 역할 배정.", "정한 Domain 에 Primary · Research · Growth 를 배정한다. Primary 만 필수다."),
        ("Level.", "처음에는 1(drafts)로 둔다. Settings 에서 언제든 올릴 수 있고, 올릴수록 Claude 가 쓸 수 있는 경로와 Skill 이 늘어난다."),
    ])
    D.h(2, "2.3 만들어지는 것")
    D.p("Vault 폴더에 `myelin.yaml`(설정), `CLAUDE.md`(Claude 운영 규칙), `10_Domains/<ID>/_MOC.md`(각 Domain 의 진입 노트), `90_System/templates`(노트 유형 9종 템플릿), `90_System/taxonomy.md`(용어 사전), `.claude/settings.json`(Claude Code 2차 방어 규칙), `.obsidian/`(Emboss 테마 · Graph View 설정)이 생긴다. git 저장소가 초기화되어 `[myelin:init]` 커밋이 남고, `90_System/hooks/pre-commit` 이 `.git/hooks/` 에 설치된다. CLI 로 같은 일을 하려면 `myelin init --vault <경로>` 를 쓴다(`--dry-run` 으로 미리 볼 수 있다).")

    # 3
    D.h(1, "3. Dashboard — 지식 신경망의 현재")
    D.p("Dashboard 는 읽기 전용이다(예외는 Obsidian 카드의 설정 쓰기 버튼 하나). 네 개의 KPI 타일, Myelination(강도 분포) 카드, 지표 표, Obsidian Graph View 카드, 최근 git 커밋을 보여 준다. 숫자는 상태 색으로 구분된다. verified 는 라임, AI 관련은 시안, 차단은 마젠타 계열이다.")
    D.h(2, "3.1 KPI 타일과 지표 9종")
    D.p("지표가 목표를 벗어나면 색으로만 알리지 않고 **달성 / 목표 이탈** 배지를 글자로 함께 보여 준다. v2.0 에서 재사용 비율 · 재사용 지연 · 산출물당 출처 세 지표가 더해졌다. 셋은 “승격한 지식이 실제로 다시 쓰이는가” 를 재는 복리 지표이며, 산출물이 하나도 없으면 `-` 로 표시된다.")
    D.table("kpi")
    D.h(2, "3.2 Myelination — 미엘린 강도")
    D.p("verified 노트가 다른 노트에서 링크되거나 `/compose` 산출물의 재료로 쓰이면 강도가 오른다. Dashboard 의 큰 숫자는 “verified 중 재사용된 비율” 이고, 옆 표는 강도별 노트 수이다. 강도는 My Domains 노트 표에도 `myelin ×N` 텍스트로 보인다. 링의 색만으로 구분하지 않는다.")
    D.table("strength")
    D.h(2, "3.3 Obsidian Graph View 카드 (v2.0)")
    D.p("Vault 의 `.obsidian/` 에 Graph View 설정이 있는지, Domain 색 그룹이 몇 개인지, 테마가 설정됐는지, Obsidian 이 이 Vault 를 등록했는지를 칩으로 보여 준다. **Obsidian 에서 Vault 열기** 버튼은 `obsidian://open?path=` 로 Obsidian 을 연다. 처음 한 번은 이 링크로 열어야 Obsidian 이 Vault 를 기억한다. Advanced URI 커뮤니티 플러그인이 있으면 같은 버튼이 Graph View 를 바로 연다. 설정이 빠져 있으면 **Graph View 설정 쓰기** 버튼이 보이고, 누르면 commit 하나로 설정 파일을 보탠다(있는 파일은 덮어쓰지 않는다).")
    D.p("주 1회 Claude Code 에서 `/myelin:review` 를 실행하면 같은 지표를 해석한 주간 보고서가 `50_Journal` 에 초안으로 생긴다.")

    # 4
    D.h(1, "4. My Domains — Domain Map 과 MOC")
    D.p("My Domains 는 설정된 Domain 을 원형 Orb 로 보여 준다. Orb 의 테두리 색은 Pack 이 정한 accent 이며 Obsidian Graph View 의 색 그룹과 같은 토큰이다. Orb 를 누르면 그 Domain 의 역할, 설명, Sub-MOC, verified 수, 노트 표가 펼쳐진다.")
    D.figure("domains", "My Domains — Domain Orb · Sub-MOC 주제 버튼 · 노트 표(강도 myelin ×N) · Obsidian 에서 _MOC 열기", max_h_cm=12)
    D.h(2, "4.1 MOC 와 Sub-MOC")
    D.p("각 Domain 폴더의 `_MOC.md` 는 그 Domain 의 지도이다. Pack 이 선언한 Sub-MOC(예: D3 의 Locomotion, Sim2Real, VLA, ROS2, HW Platform, Experiment Log)가 버튼으로 나열되고, 버튼을 누르면 그 주제의 노트를 검색해 보여 준다. **Obsidian 에서 _MOC 열기** 는 `obsidian://` 링크로 Obsidian 을 연다. MOC 노트가 아직 없으면 “Obsidian 에서 새로 만든다” 고 안내한다. MOC 는 사람 승인 영역이라 Claude 가 직접 고치지 않으며, `/myelin:link` 가 `_patches` 에 갱신안을 제안하면 Review Gate 에서 적용한다.")
    D.h(2, "4.2 노트 표")
    D.p("표의 status 칩과 강도 라벨은 아래 색 의미를 따른다. 상태는 색만으로 전달하지 않고 항상 글자 라벨을 함께 둔다. 제목을 누르면 Obsidian 에서 그 노트가 열린다.")
    D.table("chips")
    D.h(2, "4.3 Domain 을 하나 더 추가하려면")
    D.p("3개월쯤 운영해 지표에 근거가 생기면 명령줄에서 `myelin add-domain <ID>` 를 실행한다. Pack 이 선언한 Domain 이 `domains.extra` 로 추가되고 MOC·템플릿·Graph View 색 그룹이 생기며, `[myelin:add-domain]` 커밋이 남는다. 기존 노트는 바이트 하나 바뀌지 않는다. 직접 정의한 Domain 을 더하려면 먼저 Domain Packs 에서 개인 Pack 에 선언한다(6장).")

    # 5
    D.h(1, "5. Inbox / Review Gate — 승격은 읽은 사람의 클릭으로")
    D.p("Review Gate 는 MYELIN 에서 Vault 를 바꾸는 중심 화면이다. 세 구역으로 나뉜다. AI 초안, 패치 제안, Commons 후보. 어떤 버튼이든 누르면 그 한 번의 클릭이 정확히 하나의 git commit 이 된다.")
    D.figure("review", "Inbox / Review Gate — AI 초안 (내용 읽기 → Promote / Reject) · 패치 제안 · Commons 후보", max_h_cm=12)
    D.h(2, "5.1 AI 초안 (ai-drafts)")
    D.p("Claude Code 의 `/myelin:ingest`, `/myelin:distill`, `/myelin:interview` 와 MCP `write_draft` 가 만든 초안이 여기에 모인다. 각 행에는 Cyan 의 ai-draft 칩, 제목, 유형·Domain·생성일이 보인다.")
    D.bullets([
        ("내용 읽기 (v2.0).", "본문이 펼쳐진다. **이 버튼을 누르기 전에는 Promote 가 비활성**이다 — 읽지 않은 승격을 장치가 막는다(정본 Rev 2.1 §7.4)."),
        ("Promote.", "노트가 `10_Domains/<Domain>/` 으로 옮겨지고 status 가 verified, reviewed 가 오늘, origin 이 claude 에서 mixed(사람이 검토했음)로 바뀐다. 승격 순간 라임 링이 감기는 myelinated 애니메이션이 뜬다. 같은 이름의 노트가 이미 있으면 409 로 막는다."),
        ("Reject.", "`99_Archive/rejected/` 로 옮겨지고 deprecated 가 된다."),
    ])
    D.p("frontmatter 가 깨진 초안은 invalid 배지가 붙고 Promote 가 비활성이 된다. Obsidian 에서 고친 뒤 새로고침하면 된다. Domain 이 없는 초안은 승격할 Domain 을 고르라는 오류가 난다.")
    D.figure("promoted", "승격 직후 — 라임 링이 감기는 myelinated 표시와 [myelin:review] promote 커밋", max_h_cm=11)
    D.table("review")
    D.h(2, "5.2 패치 제안 (_patches)")
    D.p("`/myelin:link` 는 `10_Domains` 의 노트를 직접 고치지 않는다. 대신 `_patches/` 에 `*.patch.md` 제안 파일을 만든다. 제안은 append(끝에, 또는 `under` 로 지정한 제목 구획 아래에 덧붙임), replace(정확히 한 번 나오는 문장을 바꿈), frontmatter `## set`(허용 키 `tags` · `source` · `projects` · `aliases` · `confidence` 만) 세 종류이다.")
    D.p("v2.0 의 적용 절차는 **미리보기 → Apply** 다. 미리보기는 대상 노트에 적용했을 때의 diff, fail-closed 검사 결과(40_IP 나 confidential 노트 대상, verified 노트에 미검증 AI 노트 링크 추가, 결과 frontmatter 깨짐 등은 거부), 그리고 digest 를 보여 준다. Apply 는 그 digest 가 일치할 때만 적용되며(그 사이 패치 파일이 바뀌었으면 409), 적용된 패치는 `_patches/applied/`, 거부된 패치는 `_patches/rejected/` 로 옮겨져 이력이 남는다.")
    D.h(2, "5.3 Commons 후보")
    D.p("공개하고 싶은 verified 노트에 Obsidian 에서 `share: true` 를 적으면 이 구역에 나타난다. Hub 는 먼저 하드 게이트(verified, classification 이 internal 이하, 40_IP 밖, 기밀 노트 링크 없음)를 확인하고, 다음으로 비식별 검사(IP·이메일·전화·주민번호·시크릿·내부 URL·기밀 표시어·`deid.yaml` 금칙어)를 돌린다. 둘 다 통과한 노트만 publishable 칩이 붙고, 사용자가 “고객명·IP·기밀이 없음을 확인했습니다” 에 체크해야 **Commons 제안** 버튼이 켜진다.")
    D.bullets([
        ("push 하고 PR 을 만든다 (v2.0).", "체크하면 `publish/<노트 id>` 브랜치 커밋 뒤 `commons.remote` 로 push 하고 `gh` 로 PR 까지 만든다. 체크하지 않으면 로컬 브랜치까지만 만들고 다음에 실행할 명령을 안내한다. push 나 PR 이 실패해도 커밋은 보존된다."),
        ("두 정체성.", "Commons 노트에는 작성자 이름(author)과 원 노트 id 가 **실명**으로 남아 기여가 평판이 된다(Commons 집계의 `top_contributors`, 노트별 인용 수). 반면 Settings 의 익명 지표(opt-in)는 owner 해시와 개수만 보내며 실명·제목·본문은 포함하지 않는다."),
        ("설정 필요.", "Settings › Commons 에 로컬 clone(`commons.local`)이 없으면 제안 버튼이 비활성이고 안내가 뜬다."),
    ])
    D.p("고객명이나 IP 가 남아 있으면 blocked 로 표시되고 원인 줄이 함께 보이므로 노트를 고친 뒤 다시 시도한다. 이후 챔피언이 PR 을 검토해 합친다.")

    # 6
    D.h(1, "6. Domain Packs — Domain 은 코드가 아니라 선언")
    D.p("Pack 은 Domain 을 선언하는 YAML 이다. Core Pack(`myelin.core`)은 노트 유형 9종의 템플릿과 공통 용어를 공급하고, Reference Pack(`brian.physical-ai`)은 45년차 엔지니어의 지식 지형도(D1 AI/ML/AX Platform · D2 Embedded SW · D3 자율학습 Robotics · D4 Auto Design)와 인터뷰 질문 은행을, Org Pack(`a3.security-consulting`, `a3.incident-response`)은 사내 업무 표준을 공급한다. 개인 Pack(`me.*`)은 온보딩에서 직접 정의했거나 여기서 만든 것이다.")
    D.figure("packs", "Domain Packs — Registry 카드(Org · Reference · Personal · Core)와 개인 Pack 편집기", max_h_cm=19)
    D.h(2, "6.1 카드")
    D.p("카드에는 Pack 이름, 종류와 버전, 활성 여부(이 Vault 의 `myelin.yaml` 에 들어 있는지), 선언한 Domain 의 작은 Orb 가 보인다. 편집 버튼은 개인 Pack 에만 켜지고, Org/Reference Pack 은 보기만 된다.")
    D.h(2, "6.2 개인 Pack 만들기")
    D.bullets([
        "**새 개인 Pack** 을 누르면 편집기에 템플릿이 채워진다. pack 이름은 `me.` 로 시작해야 한다.",
        "**개인 Pack 으로 복사 (v2.0)** — Reference/Org Pack 카드에서 누르면 그 Pack 의 선언과 사이드카(taxonomy·templates)가 `packs/me.<이름>/` 으로 복사되어 편집할 수 있게 된다. 원본은 바뀌지 않는다.",
        "Domain 의 `id`(대문자·숫자·하이픈), `name`, `accent`(violet/lime/cyan/magenta), `sub_mocs`(주제 3개 안팎), `note_types`, `classification_default` 를 적는다.",
        "**검증** 으로 문법과 규칙을 확인한 뒤 **저장 (commit)** 을 누른다. Vault 의 `packs/<이름>/pack.yaml` 로 저장되고 Pack 디렉터리 전체가 커밋된다.",
        "명령줄에서 `myelin add-domain <ID>` 로 Domain 을 활성화한다. 3개월 운영 후 쓸 만하면 `myelin pack-propose packs/me.<이름>` 으로 Org Pack 후보로 제안할 수 있다.",
    ], numbered=True)

    # 7
    D.h(1, "7. Settings — Level · MCP · Commons")
    D.p("Settings 는 `myelin.yaml` 을 사람이 바꾸는 화면이다. 저장은 git commit 으로 기록되고, `CLAUDE.md` 의 자동 생성 블록이 새 Level 에 맞게 다시 쓰인다. 사용자가 `CLAUDE.md` 에 적은 개인 규칙 블록은 보존된다.")
    D.figure("settings", "Settings — Owner · Level · MCP · Commons 카드 · Skill 개방 표", max_h_cm=17)
    D.h(2, "7.1 Level — 단계적 개방")
    D.p("Level 은 Claude 에게 얼마나 많은 자율을 줄지 정하는 숫자이다. 숫자가 올라갈수록 열리는 Skill 과 쓰기 경로가 누적된다. Guard 는 이 Level 을 읽어 그 밖의 경로를 차단하므로, Level 을 올리는 것은 곧 Claude 가 쓸 수 있는 폴더를 늘리는 결정이다. 아래 Skill 개방 표가 화면에도 그대로 보인다.")
    D.table("levels")
    D.h(2, "7.2 MCP 스위치")
    D.bullets([
        ("write_draft 도구 등록.", "켜면 Claude Desktop/Code 의 MCP 가 `00_Inbox/ai-drafts` 에 새 초안을 만들 수 있다(Level 1 이상 필요). 기본은 꺼짐이며 MCP 는 읽기 전용이다."),
        ("confidential 노출.", "켜면 MCP 검색·조회가 confidential 노트와 `40_IP` 의 본문을 돌려준다. **v2.0 부터 이 스위치는 로컬 프로세스(`MYELIN_LLM=local`)에서만 효력이 있고**, 외부 세션에서는 경고와 함께 무시된다(ADR-0011)."),
    ])
    D.h(2, "7.3 Commons 카드 (v2.0)")
    D.bullets([
        ("로컬 clone 경로 (commons.local).", "`myelin commons init <path>` 로 만든 Commons 저장소의 로컬 경로. 없으면 Review Gate 의 Commons 제안이 비활성이다."),
        ("원격 저장소 (commons.remote).", "연동을 켤 때 필수. push 와 PR 의 대상이 된다."),
        ("Commons 연동 켜기.", "켜도 push 는 여전히 Review Gate 에서 사람이 체크해야 일어난다. 원격이 비어 있으면 저장이 거부된다."),
        ("익명 지표 공유 opt-in.", "정본 §11.2 — owner 해시와 개수·id 목록만 보내며 실명·제목·본문은 없다. Commons 노트의 실명 author 와는 별개다. 기본 꺼짐."),
    ])

    # 8
    D.h(1, "8. 도우미 — 우측 하단의 도움말 챗봇 (v2.0)")
    D.p("Hub 모든 화면의 우측 하단 **도우미** 말풍선을 누르면 MYELIN 도우미(AI)가 모달로 열린다. MYELIN 의 개념·사용법·절차와 이 Vault 의 상태(Level · Domain · KPI 이탈 · 검토 대기 · 관련 노트)를 답하고, 근거가 된 노트를 Obsidian 링크로 준다. 도우미는 **안내만** 한다 — 노트를 쓰거나 승격·적용·공개·커밋을 대신 하지 않으며, 대화는 서버에 저장하지 않는다(ADR-0013).")
    D.figure("assistant", "도우미 — 질문과 답, 근거 노트(Obsidian 링크), 지식 기반 모드 칩", max_h_cm=12)
    D.table("assistant")
    D.p("로컬 Ollama 가 꺼져 있으면 헤더에 “지식 기반 모드” 칩이 보이고, 내장 지식 절을 그대로 답한다. `ollama serve` 로 켜면 다음 질문부터 LLM 모드가 된다. 도우미는 기밀 라우팅 규칙을 따른다 — Ollama 주소가 loopback 이 아니면 confidential 노트와 `40_IP` 는 컨텍스트에서 뺀다.")

    # 9
    D.h(1, "9. 매일·매주 — 지식 경로를 굵게 만드는 습관")
    D.bullets([
        ("메모는 00_Inbox 에.", "회의·독서·실험에서 나온 생각을 Obsidian 으로 바로 적는다. 7일 안에 처리하는 것이 목표다(inbox_dwell_days ≤ 7)."),
        ("외부 자료는 /myelin:ingest.", "Vault 폴더에서 `claude` 를 실행하고 `/myelin:ingest <파일 또는 URL>` 을 부르면 출처·핵심 주장·인용 위치가 정리된 literature 초안이 ai-drafts 에 생긴다."),
        ("Inbox 는 /myelin:distill.", "메모 하나를 재사용 가능한 원자 노트 몇 개로 쪼갠다. 중복이 있으면 새로 만들지 않고 기존 노트 보강을 제안한다."),
        ("Review Gate 에서 승격.", "초안을 **읽고** Promote 또는 Reject. 이 클릭이 verified 를 만든다."),
        ("경험은 /myelin:interview.", "Pack 의 질문 은행에서 한 질문을 골라 짧은 인터뷰를 하고, 답을 1인칭 lesson 초안으로 남긴다."),
        ("주 1회 /myelin:review.", "지표 9종과 doctor 결과, 환류 대기 산출물을 해석한 주간 보고서가 `50_Journal` 에 생긴다. 고아 노트는 `/myelin:link` 로 연결한다."),
        ("쓸 때는 /myelin:compose.", "verified 노트만 재료로 설계서·제안서·보고서 초안을 `80_Outputs/drafts` 에 만든다. Vault 가 실제 일에 재투입되는 순간이며, 재료로 쓰인 노트의 강도가 오른다."),
        ("환류 (v2.0).", "산출물에는 `## 환류 후보` 절(재료 노트에 없어서 새로 정리한 주장·교훈·결정)이 생긴다. `/myelin:distill 80_Outputs/drafts/<산출물>.md` 로 그 절을 lesson · evergreen · adr 초안(`tags: [refeed]`)으로 되돌리고 Review Gate 에서 승격한다. 산출물 안에만 남은 지식은 다시 쓰이지 않는다."),
        ("모르면 도우미에게.", "Hub 우측 하단 말풍선. 승격·적용·공개는 여전히 사람이 버튼으로 한다."),
    ], numbered=True)
    D.table("skills")

    # 10
    D.h(1, "10. Obsidian 과 Claude Code 에서")
    D.h(2, "10.1 Obsidian — Emboss 테마 적용")
    D.p("MYELIN 은 Obsidian 의 Markdown 파일을 정본으로 삼는다. 노트를 읽고 쓰는 일은 Obsidian 에서 하고, Hub 는 그 결과를 보여 주고 승격한다. Emboss 테마는 Hub 와 같은 토큰(`design/tokens.json`)에서 생성된 Obsidian 테마이며, 백색 엠보 질감 위에 verified 는 라임, AI 초안은 시안, 차단은 마젠타로만 색이 켜진다. MYELIN 과 Obsidian 의 연동은 Obsidian 플러그인이 아니라 **테마(CSS)와 파일 규약**이다.")
    D.h(3, "먼저 — 어느 Vault 를 열었나")
    D.p("MYELIN Vault 는 `myelin.yaml` 이 있는 폴더이다. Obsidian 을 처음 실행하면 Obsidian 이 스스로 `~/Obsidian Vault` 같은 기본 Vault 를 만들어 열 수 있는데, 이 폴더는 MYELIN 과 무관하다. 창 왼쪽 아래의 Vault 이름을 눌러 MYELIN Vault 로 전환하거나, *Open folder as vault* 로 `myelin.yaml` 이 있는 폴더를 지정한다. 가장 쉬운 길은 Hub Dashboard 의 **Obsidian 에서 Vault 열기** 버튼이나 `myelin obsidian setup --open` 이다 — Obsidian 이 Vault 를 등록하고 연다.")
    D.h(3, "설치 — 테마 파일 넣기")
    D.p("`myelin init` 과 Hub 온보딩이 만드는 Vault 에는 `.obsidian/themes/MYELIN Emboss/` 와 `.obsidian/appearance.json`(라이트 모드 + MYELIN Emboss 선선택)이 이미 들어 있다. 그런 Vault 는 Obsidian 에서 열기만 하면 테마가 켜진다. 기존 Vault 는 `scripts/install.sh --vault <Vault>` 가 테마를 복사한다(같은 파일이 있으면 `.bak` 백업). 손으로 넣으려면 저장소의 `theme/obsidian-myelin-emboss/` 를 `<Vault>/.obsidian/themes/MYELIN Emboss/` 로 복사한다. 폴더 이름은 `manifest.json` 의 name(MYELIN Emboss)과 같아야 Obsidian 이 목록에 올린다.")
    D.h(3, "적용 — Settings 에서 고르기")
    D.bullets([
        "Settings(톱니) › Appearance 로 간다.",
        "Base color scheme 을 Light 로 둔다. Emboss 는 라이트 전용이며, Dark 를 고르면 깨지지 않도록 라이트 값이 그대로 적용된다.",
        "Themes 항목의 Manage 를 누르고 목록에서 MYELIN Emboss 를 고른다. 설치형 테마이므로 커뮤니티 테마 다운로드 없이 바로 보인다.",
        "목록에 없으면 Ctrl+P(명령 팔레트) › Reload app without saving 을 실행하거나 Obsidian 을 다시 연다. 그래도 없으면 폴더 이름과 manifest.json 의 name 이 같은지 확인한다.",
    ], numbered=True)
    D.h(3, "노트에 상태 표시하기")
    D.p("Obsidian 은 frontmatter 의 `cssclasses` 값을 그 노트의 보기 영역에 CSS 클래스로 붙인다. Emboss 는 이 클래스를 읽어 노트 제목에 링과 라벨을 그린다. 상태는 색만으로 전달하지 않도록 제목 옆에 글자 칩(verified / ai-draft / blocked)이 함께 붙는다.")
    D.code("---\nid: 20261018-0900-wcet-budget\ntype: evergreen\nstatus: verified\ncssclasses: [mx-status-verified]\n---")
    D.table("cssclasses")
    D.p("`cssclasses` 는 Frontmatter v2 의 12개 정본 필드가 아니지만, MYELIN 은 미지 필드를 허용·보존하므로(ADR-0003) doctor 는 이를 오류로 보지 않고 info 로만 알린다.")
    D.figure("obs-verified", "Obsidian + MYELIN Emboss — verified 노트 (제목에 라임 링 + verified 칩)", max_h_cm=10)
    D.figure("obs-ai-draft", "Obsidian + MYELIN Emboss — ai-draft 노트 (제목에 시안 glow + ai-draft 칩)", max_h_cm=10)
    D.h(2, "10.2 Obsidian — Graph View (v2.0)")
    D.p("v1.3.1 부터 Vault 의 `.obsidian/` 에는 테마 외에 세 개의 설정 파일이 함께 들어간다. Graph View 를 켜고, **Domain 폴더마다 그 Domain 의 Emboss accent 색 그룹**을 두어, Obsidian 의 링크 그래프에서 Hub 의 Domain Orb 와 같은 색으로 노트가 묶여 보이게 한다. Obsidian 에서 Ctrl/Cmd+G 또는 왼쪽 리본의 그래프 아이콘으로 연다. 노트 하나의 주변만 보려면 오른쪽 패널의 Backlinks · Outgoing links, 또는 명령 “Open local graph”.")
    D.table("obsidian")
    D.bullets([
        ("기존 Vault 에 적용.", "`myelin obsidian setup --vault <Vault> [--open]` 또는 Hub Dashboard 의 **Graph View 설정 쓰기** 버튼. 있는 파일은 덮어쓰지 않고 빠진 항목만 보태며, 읽을 수 없는 JSON 이면 아무것도 쓰지 않는다(fail-closed). `--open` 은 Obsidian 을 열어 Vault 를 등록한다."),
        ("Vault 등록.", "Obsidian 은 자기 목록에 있는 Vault 만 URI 로 연다. 처음 한 번은 Dashboard 버튼이나 `--open`, 또는 *Open folder as vault* 로 연다. 상태는 `myelin obsidian status` 나 Dashboard 카드의 “Obsidian 에 등록됨/미등록” 칩으로 본다."),
        ("색 바꾸기.", "테마와 그래프 색은 모두 `design/tokens.json` 에서 나온다. `make emboss` 로 다시 만든 뒤 `myelin obsidian setup` 을 재실행하되, 이미 있는 그룹의 색은 사용자의 것이므로 바꾸지 않는다 — 새 색을 쓰려면 `graph.json` 의 해당 그룹을 지우고 다시 실행한다."),
    ])
    D.h(3, "무엇이 달라 보이나")
    D.bullets([
        "배경과 카드가 같은 색(surface)이고 경계는 그림자로만 구분된다. 입력창은 함몰(inset), 콜아웃은 돌출(raised) 질감이다.",
        "링크·강조·태그는 Violet, 승격된 노트의 제목은 Lime 링, AI 초안은 Cyan glow 로만 색이 켜진다.",
        "테마는 생성물이다. 색을 바꾸고 싶으면 `theme.css` 가 아니라 `design/tokens.json` 을 고치고 `make emboss` 로 다시 만든 뒤 설치 스크립트로 복사한다.",
    ])
    D.figure("emboss-top", "MYELIN Emboss 디자인 시스템 — 토큰 · 카드 · 버튼 · 입력창 · 칩", max_h_cm=11)
    D.figure("emboss-bottom", "MYELIN Emboss — Domain Orb · Review Row · 표", max_h_cm=11)
    D.h(2, "10.3 Claude Code")
    D.p("Vault 폴더에서 `claude` 를 실행한다. plugin 이 설치되어 있으면 `/myelin:ingest` 같은 Skill 이 보인다(없으면 `claude --plugin-dir <repo>/plugin`). Guard 가 사람 승인 영역(`10_Domains`, `40_IP`, `90_System`, `myelin.yaml`, `CLAUDE.md`) 쓰기를 차단하므로, 차단 메시지가 뜨면 우회하지 말고 `_patches` 제안(Level 2 이상)으로 바꾼다. Bash 로 파일을 쓰는 우회도 같은 Guard 가 막는다. v2.0 의 Guard 는 세 가지다.")
    D.table("guard")
    D.bullets([
        ("기밀은 외부 LLM 에 가지 않는다 (장치).", "세션에 `MYELIN_LLM=local` 이 없으면 Claude Code 는 `40_IP/**` 와 `classification: confidential` 노트를 Read/Grep/Glob/Write 할 수 없다. 로컬 LLM 세션에서만 `MYELIN_LLM=local` 을 켠다. 외부 Claude 세션에서는 켜지 않는다."),
        ("Claude Code 밖에서 바꾼 것도 커밋 시점에 검사된다.", "`myelin init` 이 설치한 pre-commit 훅이 `myelin doctor --staged` 를 돌린다. `origin: claude` 노트가 허용 범위 밖에 있거나(ai-scope), 검토 없이 verified 이거나, confidential 인데 share:true 면 커밋이 거부된다. `git commit --no-verify` 는 Guard 가 막는다. 기존 Vault 는 `myelin doctor --install-hook`."),
        ("AI 커밋 접두사.", "Skill 이 만드는 Vault 커밋은 `[myelin:<skill>]` 로 시작한다. Hub 의 커밋도 같은 규칙이다(`[myelin:review]`, `[myelin:settings]` …)."),
    ])
    D.h(2, "10.4 MCP")
    D.p("Claude Desktop 이나 Claude Code 에 `myelin-mcp` 를 등록하면 아래 도구로 Vault 를 읽을 수 있다. 기본은 읽기 전용이고 confidential 노트의 본문은 돌려주지 않는다. 설정 방법은 `docs/guides/mcp-setup.md` 에 있다.")
    D.code("claude mcp add myelin --scope user -- uv --directory /path/to/myelin run myelin-mcp --vault /path/to/Vault\nclaude mcp list        # myelin: connected")
    D.table("mcp")

    # 11
    D.h(1, "11. 부록")
    D.h(2, "11.1 Vault 폴더 구조")
    D.table("folders")
    D.h(2, "11.2 명령줄 요약")
    D.table("cli")
    D.h(2, "11.3 문제 해결")
    D.table("trouble")
    D.h(2, "11.4 알아 두어야 할 제약 (제품 v1.3.1)")
    D.bullets([
        "다크 모드는 지원하지 않는다(토큰 `modes` 만 예약). Hub 는 localhost 전용이며 팀 모드 인증은 다음 버전이다. 공유 호스트에서 Hub 를 띄우지 않는다.",
        "Bash Guard 의 범위는 Vault 쓰기 정책이다. Vault 밖 스크립트·패키지 설치·네트워크 송신(`curl -T`, `scp`)은 검사하지 않는다. `bash/sh/perl/ruby/curl/wget` 은 Claude Code 가 확인을 묻는다. 커밋 이력으로 사후 감사.",
        "비식별 검사는 휴리스틱이다. 별명이나 프로젝트 코드는 `90_System/deid.yaml` 에 등록해야 걸린다. 챔피언의 PR 검토가 두 번째 통제이다.",
        "MCP 기밀 노출은 소유자 opt-in 이며 로컬 프로세스에서만 유효하다. 잔여 위험(40_IP 밖 confidential 노트의 제목, `.myelin/index.db` 본문 캐시, Vault 밖 스크립트)은 ADR-0011 에 적혀 있다.",
        "Obsidian Graph View 의 Vault 등록 여부는 Obsidian 의 설정 파일을 읽어 판단하므로, 비표준 경로에 설치된 Obsidian 에서는 “등록 여부 확인 불가” 로 보일 수 있다. Advanced URI 플러그인 없이 버튼은 Vault 만 연다(Graph View 는 Ctrl/Cmd+G).",
        "도우미는 로컬 Ollama 만 쓴다. 외부 LLM 어댑터·대화 저장·스트리밍은 Backlog 이다. 지식 기반 모드의 답은 내장 지식 절을 그대로 보여 주므로 질문과 정확히 맞지 않을 수 있다.",
        "Vault `CLAUDE.md` 의 personal 블록 외 영역은 Pack 재생성 시 덮어쓴다(의도된 동작). 색인 구축은 5,000 노트에 약 12초.",
        "브랜드명 MYELIN 은 사내용 가칭이다. 대외 공개 시 결합형 명칭을 검토한다.",
    ])
    D.h(2, "11.5 더 읽을 문서")
    D.bullets([
        "정본 명세 `docs/spec/MYELIN_Rev2.0.md` · Rev 2.1 부록 `docs/spec/MYELIN_Rev2.1-addendum.md`(설계 철학 갭 보완 요구사항)",
        "Markdown 가이드 `docs/guides/user-guide.md` · Obsidian 연동 `docs/guides/obsidian-setup.md` · MCP 설정 `docs/guides/mcp-setup.md` · 관리자 가이드 `docs/guides/admin-guide.md` · 파일럿 가이드 `docs/guides/pilot-guide.md` · 60분 워크숍 `docs/guides/workshop-60min.md` · Day 1 체크리스트 `docs/guides/day1-checklist.md`",
        "설계 결정 `docs/adr/0001~0014` — 특히 0006(패치 형식) · 0009(Hub 쓰기 정책) · 0010(Commons) · 0011(기밀 라우팅 장치) · 0012(미엘린 강도·재사용 지표) · 0013(도우미) · 0014(공개 저장소와 개인 Vault 분리)",
        "변경 이력 `CHANGELOG.md` · 릴리스 노트 `docs/RELEASE_NOTES.md` · 위협 모델 `docs/security/threat-model.md`",
    ])
    D.p(f"동반 파일: {XLSX_NAME}(이 문서의 표 {len(TABLES)}개), img/*.png(이 문서의 그림 {len(FIGURES)}장). Hub 스크린샷은 `make e2e` 가 만든 `docs/reports/P6/` 의 것(2026-10-27, 제품 v1.3.1)이다.", 9, MUTED)

    out = OUT / f"MYELIN_사용자가이드_{GUIDE_VER[:4]}.docx"
    d.core_properties.title = f"MYELIN 사용자 가이드 {GUIDE_VER}"
    d.core_properties.author = "WDLAB · A3 Security"
    d.core_properties.subject = f"MYELIN Personal Knowledge OS — 제품 {PRODUCT_VER} 기준"
    d.save(out)
    return out


# ----------------------------------------------------------------------------- xlsx
def build_xlsx() -> Path:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "차례"
    thin = Side(style="thin", color=BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name=FONT, size=10, bold=True, color=INK)
    body_font = Font(name=FONT, size=10, color=INK)
    head_fill = PatternFill("solid", fgColor=SURFACE_2)
    alt_fill = PatternFill("solid", fgColor=VIOLET_50)

    ws0["A1"] = f"MYELIN 사용자 가이드 {GUIDE_VER} — 표 모음"
    ws0["A1"].font = Font(name=FONT, size=14, bold=True, color=VIOLET)
    ws0["A2"] = f"{GUIDE_DATE} · 제품 {PRODUCT_VER} 기준 · {ORG}"
    ws0["A2"].font = Font(name=FONT, size=10, color=MUTED)
    ws0["A4"], ws0["B4"], ws0["C4"] = "시트", "내용", "가이드 절"
    for c in ("A4", "B4", "C4"):
        ws0[c].font = head_font
        ws0[c].fill = head_fill
        ws0[c].border = border
    section_of = {
        "changes": "0.1", "menus": "0", "domains": "2.1", "kpi": "3.1", "strength": "3.2", "chips": "4.2",
        "review": "5.1", "levels": "7.1", "skills": "9", "cli": "11.2", "folders": "11.1", "trouble": "11.3",
        "install": "1.2", "assistant": "8", "obsidian": "10.2", "cssclasses": "10.1", "guard": "10.3", "mcp": "10.4",
    }
    r = 5
    for key, (sheet, headers, rows, widths) in TABLES.items():
        ws0.cell(row=r, column=1, value=sheet).font = body_font
        ws0.cell(row=r, column=2, value=f"{headers[0]} … ({len(rows)}행)").font = body_font
        ws0.cell(row=r, column=3, value=section_of.get(key, "")).font = body_font
        for c in range(1, 4):
            ws0.cell(row=r, column=c).border = border
        r += 1
    ws0.column_dimensions["A"].width = 30
    ws0.column_dimensions["B"].width = 36
    ws0.column_dimensions["C"].width = 10

    for key, (sheet, headers, rows, widths) in TABLES.items():
        ws = wb.create_sheet(sheet[:31])
        ws.append(headers)
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i, row in enumerate(rows):
            ws.append([v.replace("`", "").replace("**", "") for v in row])
            for c in ws[i + 2]:
                c.font = body_font
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if i % 2 == 1:
                    c.fill = alt_fill
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = max(12, w * 4.2)
        ws.freeze_panes = "A2"
    out = OUT / XLSX_NAME
    wb.save(out)
    return out


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    IMG.mkdir(exist_ok=True)
    for name, src in FIGURES.values():
        if not src.exists():
            raise SystemExit(f"missing figure source: {src}")
        shutil.copyfile(src, IMG / name)
    docx_path = build_docx()
    xlsx_path = build_xlsx()
    print(docx_path)
    print(xlsx_path)


if __name__ == "__main__":
    main()
